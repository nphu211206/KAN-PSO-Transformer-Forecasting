"""Apply repeatable, author-neutral submission formatting to the JEAS DOCX."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def set_font(run, name: str, size: float, *, bold=None, italic=None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), name)


def set_style_font(style, name: str, size: float) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style_element = style.element
    style_properties = style_element.get_or_add_rPr()
    fonts = style_properties.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), name)


def ensure_line_numbering(section) -> None:
    section_properties = section._sectPr
    existing = section_properties.find(qn("w:lnNumType"))
    if existing is not None:
        section_properties.remove(existing)
    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:start"), "1")
    line_numbers.set(qn("w:restart"), "newPage")
    line_numbers.set(qn("w:distance"), "360")
    section_properties.append(line_numbers)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))
    set_font(run, "Times New Roman", 10)


def remove_table_shading(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    for shading in list(properties.findall(qn("w:shd"))):
        properties.remove(shading)


def mark_repeating_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        properties.append(header)


def set_table_geometry(table) -> None:
    """Freeze Pandoc table geometry so it is stable across Word renderers."""
    table_element = table._tbl
    grid = table_element.find(qn("w:tblGrid"))
    if grid is None:
        return
    grid_widths = [
        int(column.get(qn("w:w"), "0"))
        for column in grid.findall(qn("w:gridCol"))
    ]
    if not grid_widths or any(width <= 0 for width in grid_widths):
        return

    table.autofit = False
    properties = table_element.tblPr

    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        properties.insert(0, width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(sum(grid_widths)))

    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "120")

    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    margins = properties.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        properties.append(margins)
    for side, value in (("top", 60), ("left", 120), ("bottom", 60), ("right", 120)):
        margin = margins.find(qn(f"w:{side}"))
        if margin is None:
            margin = OxmlElement(f"w:{side}")
            margins.append(margin)
        margin.set(qn("w:type"), "dxa")
        margin.set(qn("w:w"), str(value))

    for row in table_element.findall(qn("w:tr")):
        grid_index = 0
        for cell in row.findall(qn("w:tc")):
            cell_properties = cell.find(qn("w:tcPr"))
            if cell_properties is None:
                cell_properties = OxmlElement("w:tcPr")
                cell.insert(0, cell_properties)
            grid_span = cell_properties.find(qn("w:gridSpan"))
            span = int(grid_span.get(qn("w:val"), "1")) if grid_span is not None else 1
            cell_width = sum(grid_widths[grid_index : grid_index + span])
            width_element = cell_properties.find(qn("w:tcW"))
            if width_element is None:
                width_element = OxmlElement("w:tcW")
                cell_properties.insert(0, width_element)
            width_element.set(qn("w:type"), "dxa")
            width_element.set(qn("w:w"), str(cell_width))
            grid_index += span


def format_document(input_path: Path, output_path: Path) -> None:
    document = Document(input_path)
    normal = document.styles["Normal"]
    set_style_font(normal, "Times New Roman", 12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    heading_specs = {
        "Title": (16, True, False),
        "Heading 1": (14, True, False),
        "Heading 2": (12, True, False),
        "Heading 3": (12, True, True),
    }
    for style_name, (size, bold, italic) in heading_specs.items():
        if style_name not in document.styles:
            continue
        style = document.styles[style_name]
        set_style_font(style, "Times New Roman", size)
        style.font.bold = bold
        style.font.italic = italic
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)
        ensure_line_numbering(section)
        add_page_number(section)

    for paragraph in document.paragraphs:
        paragraph.paragraph_format.widow_control = True
        if paragraph.style.name == "Normal":
            paragraph.paragraph_format.line_spacing = 2.0
        for run in paragraph.runs:
            if paragraph.style.name == "Normal":
                set_font(run, "Times New Roman", 12)

    for table in document.tables:
        set_table_geometry(table)
        if table.rows:
            mark_repeating_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                remove_table_shading(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        set_font(
                            run,
                            "Times New Roman",
                            9,
                            bold=True if row_index == 0 else None,
                        )

    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = ""
    document.core_properties.keywords = ""
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    format_document(args.input, args.output)
    print(f"Formatted {args.output}")


if __name__ == "__main__":
    main()
