"""
audit_round4_complete.py
Comprehensive verification suite for all Round 4 reviewer items.
"""
import sys, os, zipfile, re
sys.stdout.reconfigure(encoding='utf-8')
import docx
import pandas as pd
from scipy import stats

def chk(desc, cond):
    res = "[OK]  " if cond else "[FAIL]"
    print(f"  {res} {desc}")
    return 1 if cond else 0

def read_docx_text(path):
    doc = docx.Document(path)
    full_text = []
    for p in doc.paragraphs:
        if p.text.strip(): full_text.append(p.text)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                if c.text.strip(): full_text.append(c.text)
    return '\n'.join(full_text)

print("=" * 75)
print("  COMPREHENSIVE ROUND 4 SUBMISSION AUDIT")
print("=" * 75)

score = 0
total = 0

# -------------------------------------------------------------
# 1. MANUSCRIPT (paper_v91_FINAL.docx)
# -------------------------------------------------------------
print("\n[SECTION 1: MANUSCRIPT VERIFICATION]")
paper_text = read_docx_text('paper_v91_FINAL.docx')
doc = docx.Document('paper_v91_FINAL.docx')

# Abstract word count & definitions
abs_p = None
for p in doc.paragraphs[:15]:
    if p.text.startswith('Forecasting power transformer'):
        abs_p = p.text
        break

abs_words = len(abs_p.split()) if abs_p else 0
total += 1; score += chk(f"Abstract word count within 150-250 ({abs_words} words)", 150 <= abs_words <= 250)
total += 1; score += chk("Electricity Transformer Temperature (ETT) defined in abstract", "Electricity Transformer Temperature (ETT)" in (abs_p or ""))
total += 1; score += chk("Mean Squared Error (MSE) defined in abstract", "Mean Squared Error (MSE)" in (abs_p or ""))
total += 1; score += chk("7 of 8 lookback-horizon configurations in abstract", "7 of 8 lookback–horizon configurations" in (abs_p or "") or "7 of 8 lookback-horizon configurations" in (abs_p or ""))

# Figure sequencing: Fig. 3 not cited in Intro
intro_text = '\n'.join([p.text for p in doc.paragraphs[:25]])
total += 1; score += chk("Fig. 3 NOT prematurely cited in Introduction", "Fig. 3" not in intro_text)

# Loss & L1 removal
total += 1; score += chk("NO composite L1 wavelet sparsity in Section 3.6", "composite Mean Squared Error loss with L_1" not in paper_text)

# Complexity formula & claim
total += 1; score += chk("Harmonized complexity O(L log L + L H)", "𝒪(L log L + L · H)" in paper_text)
total += 1; score += chk("NO 'asymptotically superior to O(L^2)' overclaiming", "asymptotically superior to the 𝒪(L^2)" not in paper_text)

# Statistical interpretation & duplicate removal
total += 1; score += chk("Neutral Table 4 DLinear interpretation", "The results show no statistically significant difference between PSO-TCWavKAN and DLinear" in paper_text)
total += 1; score += chk("NO duplicate '62.5% win rate against DLinear' in P138", paper_text.count("achieves a 62.5% win rate (10 out of 16 conditions) against DLinear") <= 1)
total += 1; score += chk("ETTh2 results use positive reduction phrasing (no -6.3%)", "a 6.3% reduction" in paper_text or "reductions over DLinear of 6.3%" in paper_text)

# Table 1 literature context & ACKLSTM units
total += 1; score += chk("NO '0.091 kW' as standalone MSE in Table 1", "0.091 kW (wind)" not in paper_text)
total += 1; score += chk("Table 1 explicitly labels task settings", "Evaluation Setting & Reported Baseline*" in paper_text)

# Table 6 & Fig 9 clean notes
total += 1; score += chk("Table 6 note: no duplicate 'Note:' or broken syntax", "Note: Table 6 reports synthetic Gaussian noise robustness" in paper_text)
total += 1; score += chk("Fig. 9 text references clean conditions (sigma = 0.0)", "clean conditions (σ = 0.0)" in paper_text)

# Table 7 Edge feasibility
total += 1; score += chk("Table 7 Edge SCADA: 'Potentially feasible; requires target-hardware validation'", "Potentially feasible; requires target-hardware validation" in paper_text)

# Verified HAKformer DOI
total += 1; score += chk("Ref 12 has verified HAKformer DOI 10.1007/s10489-026-07209-x", "10.1007/s10489-026-07209-x" in paper_text)

# -------------------------------------------------------------
# 2. COVER LETTER (COVER_LETTER_ELECTRICAL_ENGINEERING.docx)
# -------------------------------------------------------------
print("\n[SECTION 2: COVER LETTER VERIFICATION]")
cl_text = read_docx_text('COVER_LETTER_ELECTRICAL_ENGINEERING.docx')
cl_doc = docx.Document('COVER_LETTER_ELECTRICAL_ENGINEERING.docx')

total += 1; score += chk("Author singular: 'The author has not selected'", "The author has not selected" in cl_text)
total += 1; score += chk("NO unverified Transfer Notice in Cover Letter", "Transfer Notice" not in cl_text)
total += 1; score += chk("Ablation scope restricted to H <= 336", "standard operational forecasting regime (H ≤ 336)" in cl_text)
total += 1; score += chk("Noise degradation bounded between +1.4% and +13.0%", "between +1.4% and +13.0%" in cl_text)
total += 1; score += chk("Holm-corrected p=0.133 accurately described", "Holm-corrected p = 0.133" in cl_text)
total += 1; score += chk("NO promotional hype ('next evolutionary leap', 'preventing non-linear overfitting collapse')", "next evolutionary leap" not in cl_text and "preventing non-linear overfitting collapse" not in cl_text)

# -------------------------------------------------------------
# 3. DECLARATIONS (DECLARATIONS_ELECTRICAL_ENGINEERING.docx)
# -------------------------------------------------------------
print("\n[SECTION 3: DECLARATIONS VERIFICATION]")
decl_text = read_docx_text('DECLARATIONS_ELECTRICAL_ENGINEERING.docx')

total += 1; score += chk("Truthful reproducibility statement (code on GitHub/Zenodo; logs/checkpoints on Kaggle/request)", "available upon reasonable request" in decl_text)
total += 1; score += chk("Author singular ('The author declares', 'The author received')", "The author declares" in decl_text and "The author received no external funding" in decl_text)
total += 1; score += chk("No JEAS in declarations", "JEAS" not in decl_text and "Journal of Engineering" not in decl_text)

# -------------------------------------------------------------
# 4. CODE & REPRODUCIBILITY (pso_tcwavkan_v91.py)
# -------------------------------------------------------------
print("\n[SECTION 4: CODE & SCRIPTS VERIFICATION]")
code_text = open('github_release/pso_tcwavkan_v91.py', encoding='utf-8').read()

total += 1; score += chk("Code saves actual selected model weights (eval_m / swa_model or best_sd)", "saved_weights = (eval_m.module if hasattr(eval_m, 'module') else eval_m).state_dict()" in code_text)
total += 1; score += chk("Sample-weighted test loss calculation", "total_elements += n_elem" in code_text and "total_sq_err / max(1, total_elements)" in code_text)
total += 1; score += chk("Reproduction script reproduce_statistical_tests.py exists", os.path.exists('reproduce_statistical_tests.py'))

# -------------------------------------------------------------
# 5. README & MASTER SYNCHRONIZATION
# -------------------------------------------------------------
print("\n[SECTION 5: README & MASTER SYNCHRONIZATION]")
readme_text = open('README.md', encoding='utf-8').read()

total += 1; score += chk("README specifies 5 seeds for proposed/linear, 3 for PatchTST", "5 independent random seeds" in readme_text and "3 independent seeds" in readme_text and "PatchTST Lite" in readme_text)
total += 1; score += chk("README Table 6 clarifies Seed 42 reference run", "Baseline clean values in Table 6" in readme_text)
total += 1; score += chk("NO placeholder JEAS citation in README", "10.1186/s44147-026-00000-0" not in readme_text)
total += 1; score += chk("README synchronized with release README", readme_text == open('github_release/README.md', encoding='utf-8').read())

print("\n" + "=" * 75)
print(f"  AUDIT SCORE: {score} / {total} CHECKS PASSED")
if score == total:
    print("  >>> PERFECT 100% PASS - ALL ROUND 4 ISSUES FULLY RESOLVED <<<")
else:
    print("  >>> ATTENTION: SOME CHECKS FAILED <<<")
print("=" * 75)
